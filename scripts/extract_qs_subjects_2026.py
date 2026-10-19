"""Extract selected QS 2026 broad subject area rankings from supplied workbooks."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


SOURCE_NAME = "QS World University Rankings by Subject 2026 - Public Results v1.4 (qs.com)"
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "public" / "data"
AREAS = {
    "arts-humanities": {
        "sheet": "Arts & Humanities",
        "source": Path.home() / "Downloads" / f"{SOURCE_NAME}_19.xlsx",
        "output": OUTPUT_DIR / "qs-subject-rankings-2026-arts-humanities.json",
        "count": 553,
        "last_rank": "501-550",
        "weights": (0.6, 0.2, 0.075, 0.075, 0.05),
        "url": "https://www.topuniversities.com/university-subject-rankings/arts-humanities",
    },
    "engineering-technology": {
        "sheet": "Engineering & Technology",
        "source": Path.home() / "Downloads" / f"{SOURCE_NAME}_20.xlsx",
        "output": OUTPUT_DIR / "qs-subject-rankings-2026-engineering-technology.json",
        "count": 556,
        "last_rank": "501-550",
        "weights": (0.4, 0.3, 0.1, 0.1, 0.1),
        "url": "https://www.topuniversities.com/university-subject-rankings/engineering-technology",
    },
    "life-sciences-medicine": {
        "sheet": "Life Sciences & Medicine",
        "source": Path.home() / "Downloads" / f"{SOURCE_NAME}_21.xlsx",
        "output": OUTPUT_DIR / "qs-subject-rankings-2026-life-sciences-medicine.json",
        "count": 550,
        "last_rank": "501-550",
        "weights": (0.4, 0.1, 0.2, 0.2, 0.1),
        "url": "https://www.topuniversities.com/university-subject-rankings/life-sciences-medicine",
    },
    "natural-sciences": {
        "sheet": "Natural Sciences",
        "source": Path.home() / "Downloads" / f"{SOURCE_NAME}_17.xlsx",
        "output": OUTPUT_DIR / "qs-subject-rankings-2026-natural-sciences.json",
        "count": 554,
        "last_rank": "501-550",
        "weights": (0.4, 0.2, 0.15, 0.15, 0.1),
        "url": "https://www.topuniversities.com/university-subject-rankings/natural-sciences",
    },
    "social-sciences-management": {
        "sheet": "Social Sciences & Management",
        "source": Path.home() / "Downloads" / f"{SOURCE_NAME}_18.xlsx",
        "output": OUTPUT_DIR / "qs-subject-rankings-2026-social-sciences-management.json",
        "count": 550,
        "last_rank": "501-550",
        "weights": (0.5, 0.3, 0.075, 0.075, 0.05),
        "url": "https://www.topuniversities.com/university-subject-rankings/social-sciences-management",
    },
}
HEADERS = ("2026", "2025", "INSTITUTION", "COUNTRY/TERRITORY", "ACADEMIC", "ACADEMIC RANK", "EMPLOYER", "EMPLOYER RANK", "CITATIONS", "CITATIONS RANK", "H", "H RANK", "IRN", "IRN RANK", "SCORE")


def read_source(source: Path, area: str) -> dict:
    config = AREAS[area]
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[config["sheet"]]
    if sheet["A1"].value != "2026 QS World University Rankings by Subject" or sheet["A2"].value != config["sheet"]:
        raise ValueError("Unexpected workbook edition or area")
    if tuple(cell.value for cell in sheet[4]) != HEADERS:
        raise ValueError("Unexpected ranking columns")

    weights = tuple(sheet.cell(2, column).value for column in range(6, 11))
    if weights != config["weights"]:
        raise ValueError(f"Unexpected methodology weights: {weights}")

    universities = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), start=5):
        rank, _, university, country = row[:4]
        score = row[14]
        if not isinstance(rank, (int, str)) or not isinstance(university, str) or not isinstance(country, str):
            raise ValueError(f"Missing rank, university or country at row {row_number}")
        if score in (None, "-"):
            score = None
        elif not isinstance(score, (int, float)):
            raise ValueError(f"Unexpected score at row {row_number}: {score!r}")
        universities.append({
            "index": len(universities) + 1,
            "rank": str(rank),
            "university": university.strip(),
            "location": country.strip(),
            "score": score,
        })

    if len(universities) != config["count"] or universities[0]["rank"] != "1" or universities[-1]["rank"] != config["last_rank"]:
        raise ValueError("Unexpected number or range of ranking entries")

    return {
        "edition": 2026,
        "area": config["sheet"],
        "source_file": source.name,
        "source_url": config["url"],
        "count": len(universities),
        "methodology_weights": {"academic": weights[0], "employer": weights[1], "citations_per_paper": weights[2], "h_index": weights[3], "international_research_network": weights[4]},
        "universities": universities,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--area", choices=AREAS, default="arts-humanities")
    parser.add_argument("--source", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    source = args.source or AREAS[args.area]["source"]
    output = args.output or AREAS[args.area]["output"]
    dataset = read_source(source, args.area)
    if args.verify:
        saved = json.loads(output.read_text(encoding="utf-8"))
        if saved != dataset:
            raise ValueError("Saved JSON differs from the source workbook")
        print(f"Verified {dataset['count']} complete rows against {source.name}")
    else:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"Saved {dataset['count']} entries; {sum(item['score'] is None for item in dataset['universities'])} without a published score")


if __name__ == "__main__":
    main()
