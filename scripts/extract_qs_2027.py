"""Inspect and extract QS World University Rankings 2027 from the user-supplied XLSX."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_SOURCE = Path.home() / "Downloads" / "2027 QS World University Rankings 1.3 (For qs.com).xlsx"
DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "public" / "data" / "qs-world-university-rankings-2027-top1000.json"
EXPECTED_HEADERS = {
    0: "Index",
    1: "Rank",
    3: "Name",
    4: "Country/Territory",
    5: "Region",
    28: "Overall SCORE",
}


def inspect(source: Path) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    for sheet in workbook:
        print(f"SHEET {sheet.title!r}: {sheet.max_row} rows x {sheet.max_column} columns")
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_number in {1, 2, 3, 4, 5, 6, 1002, 1003, 1004, 1005, 1506, 1507}:
                print(row_number, repr(row))


def extract(source: Path, output: Path) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"Unexpected worksheets: {workbook.sheetnames}")

    rows = workbook.active.iter_rows(values_only=True)
    title = next(rows)
    next(rows)
    headers = next(rows)
    if title[0] != "2027 QS World University Rankings":
        raise ValueError("Unexpected ranking edition")
    for column, expected in EXPECTED_HEADERS.items():
        if headers[column] != expected:
            raise ValueError(f"Unexpected column {column + 1}: {headers[column]!r}")

    universities = []
    for index, row in enumerate(rows, start=1):
        if index > 1000:
            break
        if row[0] != index:
            raise ValueError(f"Unexpected source index at entry {index}: {row[0]!r}")
        rank, name, country, region, score = row[1], row[3], row[4], row[5], row[28]
        if not isinstance(rank, (int, str)) or not name or not country or not region:
            raise ValueError(f"Missing rank, name, country, or region at entry {index}")
        if score == "-" or score is None:
            overall_score = None
        elif isinstance(score, (int, float)):
            overall_score = score
        else:
            raise ValueError(f"Unexpected overall score at entry {index}: {score!r}")
        universities.append(
            {
                "index": index,
                "rank": str(rank),
                "university": str(name).strip(),
                "location": str(country).strip(),
                "region": str(region).strip(),
                "overall_score": overall_score,
            }
        )

    if len(universities) != 1000:
        raise ValueError(f"Expected 1000 entries, got {len(universities)}")

    dataset = {
        "edition": 2027,
        "source_file": source.name,
        "source_url": "https://www.topuniversities.com/world-university-rankings",
        "data_copyright_url": "https://www.topuniversities.com/data-copyright",
        "attribution": "Source: QS Quacquarelli Symonds (www.topuniversities.com). Copyright © 2026 QS Quacquarelli Symonds Ltd.",
        "note": "Location is the Country/Territory column. The source does not provide city. Region is copied from the source. Null overall_score means the source shows '-' or has no value.",
        "count": len(universities),
        "universities": universities,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    missing_scores = sum(item["overall_score"] is None for item in universities)
    print(f"Saved {len(universities)} entries to {output}")
    print(f"First rank: {universities[0]['rank']}; last rank: {universities[-1]['rank']}")
    print(f"Missing overall scores in source: {missing_scores}")


def verify(source: Path, output: Path) -> None:
    dataset = json.loads(output.read_text(encoding="utf-8"))
    universities = dataset["universities"]
    if dataset["edition"] != 2027 or dataset["count"] != 1000 or len(universities) != 1000:
        raise ValueError("Wrong edition or entry count in output")
    if [item["index"] for item in universities] != list(range(1, 1001)):
        raise ValueError("Output indices are incomplete or out of order")

    sheet = load_workbook(source, read_only=True, data_only=True).active
    source_rows = sheet.iter_rows(min_row=4, max_row=1003, values_only=True)
    for index, (item, source_row) in enumerate(zip(universities, source_rows), start=1):
        expected_score = source_row[28]
        if expected_score in (None, "-"):
            expected_score = None
        if (
            item["rank"] != str(source_row[1])
            or item["university"] != str(source_row[3]).strip()
            or item["location"] != str(source_row[4]).strip()
            or item["region"] != str(source_row[5]).strip()
            or item["overall_score"] != expected_score
        ):
            raise ValueError(f"Output differs from source at entry {index}")
    print("Verified 1000 ordered records and all displayed fields against the XLSX")
    print(f"Entries without published overall score: {sum(item['overall_score'] is None for item in universities)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    if args.inspect:
        inspect(args.source)
    elif args.verify:
        verify(args.source, args.output)
    else:
        extract(args.source, args.output)
